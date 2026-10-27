from pathlib import Path

import re

from docx import Document as DocxDocument
from docx.oxml.ns import qn
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph as DocxParagraph

from langchain_core.documents import Document
from langchain_community.document_loaders import PyMuPDFLoader, TextLoader
from langchain_ollama import OllamaEmbeddings
from langchain_text_splitters import RecursiveCharacterTextSplitter
from sqlalchemy import delete, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.config import BASE_DIR, settings
from app.models import Chunk, FileRecord
from app.services.settings_service import get_system_settings

SUPPORTED_TYPES = {"md", "docx", "pdf", "txt"}


def get_embeddings() -> OllamaEmbeddings:
    return OllamaEmbeddings(
        model=settings.embed_model,
        base_url=settings.ollama_base_url,
        client_kwargs={"trust_env": False},
    )


def _iter_docx_blocks(doc):
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield DocxParagraph(child, doc)
        elif child.tag == qn("w:tbl"):
            yield DocxTable(child, doc)


def _heading_level(paragraph) -> int | None:
    style_name = (paragraph.style.name or "").lower() if paragraph.style else ""
    if style_name.startswith("heading"):
        try:
            return int(style_name.split()[-1])
        except ValueError:
            return 1
    p_pr = paragraph._p.pPr
    if p_pr is not None:
        outline = p_pr.find(qn("w:outlineLvl"))
        if outline is not None:
            return int(outline.get(qn("w:val"))) + 1
    return None


def _table_to_markdown(table) -> str:
    rows = []
    for row in table.rows:
        cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
        if cells:
            rows.append(cells)
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    lines = ["| " + " | ".join(row + [""] * (width - len(row))) + " |" for row in rows]
    lines.insert(1, "| " + " | ".join(["---"] * width) + " |")
    return "\n".join(lines)


def _load_docx(path: Path) -> list:
    doc = DocxDocument(str(path))
    blocks = []
    for item in _iter_docx_blocks(doc):
        if isinstance(item, DocxParagraph):
            text = item.text.strip()
            if not text:
                continue
            level = _heading_level(item)
            if level:
                blocks.append(("heading", text, level))
            else:
                blocks.append(("paragraph", text, None))
        else:
            markdown = _table_to_markdown(item)
            if markdown:
                blocks.append(("table", markdown, None))
    return blocks


def _group_table_rows(block: str, chunk_size: int) -> list[str]:
    table_rows = _split_table_rows(block)
    grouped: list[str] = []
    current: list[str] = []
    current_size = 0
    for row in table_rows:
        if current and current_size + len(row) > chunk_size:
            grouped.append("\n\n".join(current))
            current = []
            current_size = 0
        current.append(row)
        current_size += len(row)
    if current:
        grouped.append("\n\n".join(current))
    return grouped


def _split_docx(blocks: list, chunk_size: int, chunk_overlap: int) -> list[Document]:
    splitter = RecursiveCharacterTextSplitter(chunk_size=chunk_size, chunk_overlap=chunk_overlap)
    sections: list[tuple[str, list]] = []
    heading_path: list[tuple[int, str]] = []
    current_blocks: list = []

    def flush_section():
        if current_blocks:
            path_label = " > ".join(title for _, title in heading_path)
            sections.append((path_label, list(current_blocks)))
            current_blocks.clear()

    for kind, content, level in blocks:
        if kind == "heading":
            flush_section()
            while heading_path and heading_path[-1][0] >= level:
                heading_path.pop()
            heading_path.append((level, content))
        else:
            current_blocks.append((kind, content))
    flush_section()

    documents: list[Document] = []
    for path_label, section_blocks in sections:
        if not section_blocks:
            continue
        chunks: list[str] = []
        paragraph_buffer: list[str] = []

        def flush_paragraphs():
            if paragraph_buffer:
                chunks.extend(splitter.split_text("\n\n".join(paragraph_buffer)))
                paragraph_buffer.clear()

        for kind, content in section_blocks:
            if kind == "table":
                flush_paragraphs()
                chunks.extend(_group_table_rows(content, chunk_size))
            else:
                paragraph_buffer.append(content)
        flush_paragraphs()

        prefix = f"{path_label}\n" if path_label else ""
        documents.extend(Document(page_content=prefix + chunk) for chunk in chunks)
    return documents


def _load_documents(path: Path, file_type: str):
    if file_type == "md":
        return TextLoader(str(path), encoding="utf-8").load()
    if file_type == "txt":
        return TextLoader(str(path), encoding="utf-8").load()
    if file_type == "pdf":
        return PyMuPDFLoader(str(path), extract_tables="markdown").load()
    raise ValueError(f"不支持的文件类型: {file_type}")


def _split_table_rows(block: str) -> list[str]:
    lines = [line.strip() for line in block.splitlines() if line.strip()]
    if len(lines) < 3:
        return []
    headers = [cell.strip() for cell in lines[0].strip("|").split("|")]
    if len(headers) < 2:
        return []
    rows = []
    for line in lines[2:]:
        cells = [cell.strip() for cell in line.strip("|").split("|")]
        pairs = [f"{header}: {cell}" for header, cell in zip(headers, cells) if cell]
        if pairs:
            rows.append("\n".join(pairs))
    return rows


def _split_documents(docs: list[Document], chunk_size: int, chunk_overlap: int) -> list[Document]:
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
    )
    chunks: list[Document] = []
    for doc in docs:
        for block in re.split(r"\n\s*\n", doc.page_content):
            lines = [line for line in block.splitlines() if line.strip()]
            if len(lines) >= 3 and all("|" in line for line in lines) and "---" in lines[1]:
                grouped = _group_table_rows(block, chunk_size)
                chunks.extend(Document(page_content=item, metadata=doc.metadata) for item in grouped)
            else:
                chunks.extend(splitter.split_documents([Document(page_content=block.strip(), metadata=doc.metadata)]))
    return chunks


async def ingest_file(
    db: AsyncSession, file_id: int, path: Path, file_type: str, filename: str, kb_id: int | None = None
) -> int:
    system_settings = await get_system_settings(db)
    if file_type == "docx":
        chunks = _split_docx(_load_docx(path), system_settings.chunk_size, system_settings.chunk_overlap)
    else:
        docs = _load_documents(path, file_type)
        chunks = _split_documents(docs, system_settings.chunk_size, system_settings.chunk_overlap)
    unique_chunks: list[Document] = []
    seen_contents: set[str] = set()
    for chunk in chunks:
        key = chunk.page_content.strip()
        if key and key not in seen_contents:
            seen_contents.add(key)
            unique_chunks.append(chunk)
    chunks = unique_chunks
    if not chunks:
        return 0

    texts: list[str] = []
    rows: list[Chunk] = []
    for index, chunk in enumerate(chunks):
        page_no = int(chunk.metadata.get("page") or 0)
        row = Chunk(
            file_id=file_id, chunk_index=index, page_no=page_no, content=chunk.page_content
        )
        db.add(row)
        rows.append(row)
        texts.append(chunk.page_content)
    await db.flush()
    vectors = await get_embeddings().aembed_documents(texts)
    for row, vector in zip(rows, vectors):
        row.embedding = vector
    return len(chunks)


async def process_file_record(
    file_id: int, path: Path, file_type: str, filename: str, kb_id: int | None
) -> None:
    from app.database import SessionLocal

    try:
        async with SessionLocal() as db:
            await ingest_file(db, file_id, path, file_type, filename, kb_id)
            record = await db.get(FileRecord, file_id)
            if record:
                record.status = "ready"
            await db.commit()
    except Exception as exc:
        async with SessionLocal() as db:
            await db.execute(delete(Chunk).where(Chunk.file_id == file_id))
            record = await db.get(FileRecord, file_id)
            if record:
                record.status = "failed"
                record.status_message = str(exc)[:255]
            await db.commit()


async def delete_file_vectors(db: AsyncSession, file_id: int) -> None:
    await db.execute(delete(Chunk).where(Chunk.file_id == file_id))
    await db.flush()


async def reindex_missing_embeddings(db: AsyncSession) -> int:
    """升级后缺失向量的文件按原始上传文件重建向量（启动时执行）。"""
    result = await db.execute(
        select(FileRecord.id)
        .join(Chunk, Chunk.file_id == FileRecord.id)
        .where(Chunk.embedding.is_(None))
        .distinct()
    )
    file_ids = result.scalars().all()
    if not file_ids:
        return 0

    upload_dir = BASE_DIR / settings.upload_dir
    rebuilt = 0
    for file_id in file_ids:
        record = await db.get(FileRecord, file_id)
        if record is None:
            continue
        matches = list(upload_dir.glob(f"{file_id}_*"))
        if not matches:
            continue
        await delete_file_vectors(db, file_id)
        rebuilt += await ingest_file(
            db, file_id, matches[0], record.file_type, record.filename, record.kb_id
        )
    await db.commit()
    return rebuilt
