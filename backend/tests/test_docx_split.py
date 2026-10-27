from docx import Document as DocxDocument

from app.services.document_service import _load_docx, _split_docx


def _build_docx(tmp_path):
    doc = DocxDocument()
    doc.add_heading("使用指南", level=1)
    doc.add_paragraph("打开 Cherry Studio → 设置 → MCP 服务器")
    doc.add_paragraph("按提示完成登录。")
    doc.add_heading("MCP 配置", level=2)
    doc.add_paragraph("配置服务器 URL 后保存。")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "配置项"
    table.cell(0, 1).text = "值"
    table.cell(1, 0).text = "服务器名称"
    table.cell(1, 1).text = "基础信息"
    path = tmp_path / "test.docx"
    doc.save(str(path))
    return path


def test_docx_split_merges_short_paragraphs(tmp_path):
    chunks = _split_docx(_load_docx(_build_docx(tmp_path)), 800, 100)
    contents = [chunk.page_content for chunk in chunks]

    assert "使用指南" not in contents
    assert any("打开 Cherry Studio" in content and "按提示完成登录" in content for content in contents)


def test_docx_split_adds_breadcrumb(tmp_path):
    chunks = _split_docx(_load_docx(_build_docx(tmp_path)), 800, 100)
    contents = [chunk.page_content for chunk in chunks]

    assert any(content.startswith("使用指南 > MCP 配置\n") for content in contents)


def test_docx_split_expands_table_rows(tmp_path):
    chunks = _split_docx(_load_docx(_build_docx(tmp_path)), 800, 100)
    contents = [chunk.page_content for chunk in chunks]

    assert any("配置项: 服务器名称" in content for content in contents)
